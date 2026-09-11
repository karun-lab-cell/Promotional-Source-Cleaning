import streamlit as st
import pandas as pd
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


# =========================================================
# PAGE CONFIG
# =========================================================

st.set_page_config(
    page_title="Promotional Source Cleaner",
    page_icon="📊",
    layout="wide"
)

st.title("📊 Promotional Source Cleaner")
st.write(
    "Upload your raw publisher file, clean Promotional Source data, "
    "validate the results, and download the formatted Excel file."
)


# =========================================================
# FILE UPLOAD
# =========================================================

uploaded_file = st.file_uploader(
    "📁 Upload Raw CSV / XLSX File",
    type=["csv", "xlsx"]
)


if uploaded_file is not None:

    # -----------------------------------------------------
    # IMPORT RAW FILE
    # -----------------------------------------------------

    try:

        if uploaded_file.name.lower().endswith(".csv"):
            df = pd.read_csv(uploaded_file)
        else:
            df = pd.read_excel(uploaded_file)

        st.success(
            f"File uploaded successfully: {uploaded_file.name}"
        )

    except Exception as e:

        st.error(f"Error reading file: {e}")
        st.stop()


    # -----------------------------------------------------
    # RAW DATA SUMMARY
    # -----------------------------------------------------

    st.subheader("Raw Data Summary")

    col1, col2 = st.columns(2)

    col1.metric("Rows", f"{len(df):,}")
    col2.metric("Columns", f"{len(df.columns):,}")

    st.write("**Available Columns:**")
    st.write(df.columns.tolist())

    st.subheader("Raw Data Preview")

    st.dataframe(
        df.head(10),
        use_container_width=True
    )


    # =====================================================
    # CLEAN DATA BUTTON
    # =====================================================

    if st.button(
        "🚀 CLEAN DATA",
        type="primary"
    ):

        with st.spinner("Cleaning Promotional Source data..."):

            try:

                # =================================================
                # STEP 1: CLEAN COLUMN NAMES
                # =================================================

                df.columns = df.columns.str.strip()


                # =================================================
                # STEP 2: RENAME COLUMNS
                # =================================================

                df = df.rename(columns={
                    "id": "ID",
                    "manager": "Manager",
                    "status": "Status",
                    "parent_pub_id": "Parent_pub_id",
                    "primarypromotionalmethod":
                        "Original_Promotional_Source"
                })


                # =================================================
                # STEP 3: CLEAN ID COLUMNS
                # =================================================

                df["Parent_pub_id"] = (
                    df["Parent_pub_id"]
                    .fillna("")
                    .astype(str)
                    .str.replace(r"\.0$", "", regex=True)
                    .str.strip()
                )

                df["ID"] = (
                    df["ID"]
                    .fillna("")
                    .astype(str)
                    .str.replace(r"\.0$", "", regex=True)
                    .str.strip()
                )


                # =================================================
                # STEP 4: CREATE Pub_Id_replaced
                # =================================================

                df["Pub_Id_replaced"] = df["Parent_pub_id"].where(
                    df["Parent_pub_id"] != "",
                    df["ID"]
                )


                # =================================================
                # STEP 5: CREATE PROMOTIONAL SOURCE COLUMN
                # =================================================

                df["Original_Promotional_Source"] = (
                    df["Original_Promotional_Source"]
                    .fillna("")
                    .astype(str)
                )

                df["Promotional Source"] = (
                    df["Original_Promotional_Source"]
                )


                # =================================================
                # STEP 6: BASIC PROMOTIONAL SOURCE CLEANING
                # =================================================

                df["Promotional Source"] = (
                    df["Promotional Source"]
                    .str.replace(
                        "[", "", regex=False
                    )
                    .str.replace(
                        "]", "", regex=False
                    )
                    .str.replace(
                        '"', "", regex=False
                    )
                    .str.replace(
                        ":", "-", regex=False
                    )
                    .str.replace(
                        "Buying", "buy", regex=False
                    )
                    .str.replace(
                        "Influencer", "Creator",
                        regex=False
                    )
                    .str.replace(
                        r"\s+", " ",
                        regex=True
                    )
                    .str.strip()
                )


                # =================================================
                # STEP 7: NORMALIZE HYPHEN SPACING
                # =================================================

                df["Promotional Source"] = (
                    df["Promotional Source"]
                    .str.replace(
                        r"\s*-\s*",
                        " - ",
                        regex=True
                    )
                    .str.replace(
                        r"\s+",
                        " ",
                        regex=True
                    )
                    .str.strip()
                )


                # =================================================
                # STEP 8: SINGLE SOURCE MAPPING
                # =================================================

                single_source_mapping = {

                    "Search": "Media buy - Search",
                    "Social": "Media buy - Social",
                    "Native": "Media buy - Native",
                    "Display": "Media buy - Display",
                    "Video": "Media buy - Display",
                    "Pop": "Media buy - Display",
                    "Push": "Media buy - Display",

                    "Display - Banner":
                        "Media buy - Display",

                    "Display - In-text":
                        "Media buy - Display",

                    "Dispaly - Banner":
                        "Media buy - Display",

                    "Dispaly - In-text":
                        "Media buy - Display",

                    "Email": "Database - Email",
                    "SMS": "Database - SMS",
                    "Telegram": "Database - Telegram",
                    "WhatsApp": "Database - WhatsApp",

                    "Content": "Website - Content",
                    "Coupon": "Website - Coupon",
                    "Cashback": "Website - Cashback",

                    "Mobile App":
                        "Website - Mobile App",

                    "MobileApp":
                        "Website - Mobile App",

                    "Creator": "Partners - Creator",
                    "Software": "Partners - Software",

                    "Sub-network":
                        "Partners - Sub-network",

                    "Sub - network":
                        "Partners - Sub-network",

                    "Reseller": "Partners - Reseller",
                    "Internal": "Partners - Internal"
                }


                def map_single_source(value):

                    value = value.strip()

                    if value in single_source_mapping:
                        return single_source_mapping[value]

                    return value


                df["Promotional Source"] = (
                    df["Promotional Source"]
                    .apply(map_single_source)
                )


                # =================================================
                # STEP 9: DISPLAY VARIATIONS
                # =================================================

                display_values = {

                    "Display",
                    "Video",
                    "Pop",
                    "Push",
                    "Display - Banner",
                    "Display - In-text",
                    "Dispaly - Banner",
                    "Dispaly - In-text"
                }


                df["Promotional Source"] = (
                    df["Promotional Source"]
                    .apply(
                        lambda x:
                        "Media buy - Display"
                        if x.strip() in display_values
                        else x
                    )
                )


                # =================================================
                # STEP 10: NORMALIZE SUB-NETWORK
                # =================================================

                df["Promotional Source"] = (
                    df["Promotional Source"]
                    .str.replace(
                        r"Partners\s*-\s*Sub\s*-\s*network",
                        "Partners - Sub-network",
                        regex=True
                    )
                )


                # =================================================
                # STEP 11: INCENT NORMALIZATION
                # =================================================

                df.loc[
                    df["Promotional Source"]
                    .str.strip()
                    .str.lower() == "incent",
                    "Promotional Source"
                ] = "Incent"


                # =================================================
                # STEP 12: ARRANGE FINAL COLUMNS
                # =================================================

                required_columns = [
                    "ID",
                    "Manager",
                    "Status",
                    "Pub_Id_replaced",
                    "Parent_pub_id",
                    "Promotional Source",
                    "Original_Promotional_Source"
                ]

                df = df[required_columns]


                # =================================================
                # STEP 13: CREATE EXCEL IN MEMORY
                # =================================================

                output_buffer = BytesIO()

                df.to_excel(
                    output_buffer,
                    index=False,
                    sheet_name="Data"
                )

                output_buffer.seek(0)


                # =================================================
                # STEP 14: EXCEL FORMATTING
                # =================================================

                wb = load_workbook(output_buffer)

                ws = wb.active

                # General settings
                ws.sheet_view.showGridLines = False
                ws.freeze_panes = "A2"


                # Colors / Styles
                dark_blue = "17365D"
                white = "FFFFFF"

                header_fill = PatternFill(
                    fill_type="solid",
                    fgColor=dark_blue
                )

                header_font = Font(
                    name="Verdana",
                    size=8,
                    bold=True,
                    color=white
                )

                body_font = Font(
                    name="Verdana",
                    size=8
                )

                thin_side = Side(
                    style="thin",
                    color="BFBFBF"
                )

                data_border = Border(
                    left=thin_side,
                    right=thin_side,
                    top=thin_side,
                    bottom=thin_side
                )


                # Apply body formatting
                for row in ws.iter_rows():

                    for cell in row:

                        cell.font = body_font
                        cell.border = data_border


                # Header formatting
                for cell in ws[1]:

                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = data_border


                # =================================================
                # STEP 15: COLUMN WIDTHS
                # =================================================

                # A:E auto-fit
                for col_num in range(1, 6):

                    column_letter = get_column_letter(
                        col_num
                    )

                    max_length = 0

                    for cell in ws[column_letter]:

                        if cell.value is not None:

                            cell_length = len(
                                str(cell.value)
                            )

                            if cell_length > max_length:
                                max_length = cell_length

                    ws.column_dimensions[
                        column_letter
                    ].width = max_length + 2


                # F width based on header
                ws.column_dimensions["F"].width = (
                    len(str(ws["F1"].value)) + 2
                )


                # =================================================
                # STEP 16: HIDE ORIGINAL SOURCE
                # =================================================

                # Column G
                ws.column_dimensions["G"].hidden = True


                # =================================================
                # STEP 17: FORMAT ID COLUMNS
                # =================================================

                for column in ["A", "D", "E"]:

                    for cell in ws[column][1:]:

                        if cell.value is not None:

                            try:

                                cell.value = int(
                                    float(
                                        str(
                                            cell.value
                                        ).strip()
                                    )
                                )

                            except (
                                ValueError,
                                TypeError
                            ):

                                pass

                        cell.number_format = "0"


                # =================================================
                # STEP 18: SAVE FINAL EXCEL TO MEMORY
                # =================================================

                final_buffer = BytesIO()

                wb.save(final_buffer)

                final_buffer.seek(0)


                # =================================================
                # SUCCESS / VALIDATION
                # =================================================

                st.success(
                    "✅ Promotional Source cleaning completed!"
                )

                st.subheader("Cleaning Summary")

                col1, col2, col3 = st.columns(3)

                col1.metric(
                    "Final Rows",
                    f"{len(df):,}"
                )

                col2.metric(
                    "Final Columns",
                    f"{len(df.columns):,}"
                )

                col3.metric(
                    "Unique Promotional Sources",
                    f"{df['Promotional Source'].nunique():,}"
                )


                # =================================================
                # PROMOTIONAL SOURCE SUMMARY
                # =================================================

                st.subheader(
                    "Promotional Source Distribution"
                )

                source_summary = (
                    df["Promotional Source"]
                    .value_counts()
                    .reset_index()
                )

                source_summary.columns = [
                    "Promotional Source",
                    "Count"
                ]

                st.dataframe(
                    source_summary,
                    use_container_width=True
                )


                # =================================================
                # CLEANED DATA PREVIEW
                # =================================================

                st.subheader("Cleaned Data Preview")

                st.dataframe(
                    df.head(10),
                    use_container_width=True
                )


                # =================================================
                # DOWNLOAD BUTTON
                # =================================================

                st.download_button(
                    label="⬇️ DOWNLOAD CLEANED EXCEL",
                    data=final_buffer,
                    file_name="Promotional_Source_Cleaned.xlsx",
                    mime=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    )
                )


            except Exception as e:

                st.error(
                    f"❌ Error during cleaning: {e}"
                )

                st.exception(e)
